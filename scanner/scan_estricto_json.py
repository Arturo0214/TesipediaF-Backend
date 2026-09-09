#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Escáner anti-IA (MODO ESTRICTO) — salida JSON para la web.
Reusa EXACTAMENTE los criterios de scan_report_estricto.py del proyecto
escaner-anti-ia (blacklist ampliada, construcciones simétricas, enumeraciones
mecánicas, guiones largos, aperturas encadenadas y variación de ritmo CV).
No genera PDF; imprime JSON en stdout.

Uso:  python3 scan_estricto_json.py <documento.docx>
Requiere: python-docx
Escala:  BAJO 0–19 · MEDIO 20–44 · ALTO 45–100
"""
import sys, re, json
try:
    import docx
except Exception:
    print(json.dumps({"error": "python-docx no instalado"})); sys.exit(0)

BLACKLIST = [r'en este sentido', r'cabe destacar', r'cabe mencionar', r'cabe señalar', r'cabe resaltar',
 r'es importante señalar', r'es importante mencionar', r'es importante destacar', r'desde esta óptica',
 r'desde esta perspectiva', r'en ese orden de ideas', r'resulta pertinente', r'resulta fundamental',
 r'en consonancia con', r'de igual manera', r'de igual forma', r'en el marco de', r'es fundamental comprender',
 r'juega un papel fundamental', r'juega un papel crucial', r'ha sido ampliamente estudiado',
 r'es un tema de gran relevancia', r'en la sociedad actual', r'en el mundo actual', r'diversos autores han señalado',
 r'en términos generales', r'a lo largo de la historia', r'sin lugar a dudas', r'es innegable que',
 r'resulta evidente que', r'lo anterior permite', r'en primer lugar', r'en segundo lugar', r'en tercer lugar',
 r'hoy en d[ií]a', r'en la actualidad', r'de vital importancia', r'de suma importancia', r'pilar fundamental',
 r'piedra angular', r'en aras de', r'dicho lo anterior', r'es menester', r'en definitiva', r'en resumidas cuentas',
 r'a fin de cuentas', r'dicho de otro modo', r'en s[ií]ntesis', r'no cabe duda', r'abanico de posibilidades',
 r'desempeña un papel (?:fundamental|crucial|clave|primordial)', r'un sinf[ií]n de', r'en tal sentido', r'por consiguiente']
SYMM = [r'no es\s+\w+[\w\s,]{0,40}?\s+sino\b', r'no se trata de\s+[\w\s,]{0,40}?\s+sino\b',
      r'más que\s+[\w\s,]{0,30}?\s+lo que se busca', r'lejos de\s+[\w\s,]{0,30}?\s+se propone',
      r'no s[oó]lo\s+[\w\s,]{0,60}?\s+sino tambi[eé]n\b', r'no s[oó]lo\s+[\w\s,]{0,60}?\s+sino\b',
      r'no radica\s+[\w\s,]{0,40}?\s+sino\b', r'no consiste\s+[\w\s,]{0,40}?\s+sino\b']
ENUM = [r'\bpor un lado\b[\w\s,]{0,120}\bpor otro lado\b', r'\bprimero\b[\w\s,]{0,80}\bsegundo\b[\w\s,]{0,80}\btercero\b',
      r'\ben primer lugar\b[\w\s,]{0,140}\ben segundo lugar\b', r'\bpor una parte\b[\w\s,]{0,120}\bpor otra parte\b']
STRICT = [r'en suma', r'en efecto', r'asimismo', r'por ende', r'por consiguiente', r'por lo tanto',
 r'de esta manera', r'de esta forma', r'de tal manera', r'de tal forma', r'de igual modo',
 r'en ese sentido', r'a su vez', r'por su parte', r'es decir', r'esto es,', r'en otras palabras',
 r'con base en lo anterior', r'aunado a lo anterior', r'adicionalmente', r'vale la pena',
 r'no obstante', r'sin embargo', r'por otro lado', r'por otra parte', r'por un lado', r'por una parte',
 r'finalmente', r'por [uú]ltimo', r'en conclusi[oó]n', r'para concluir', r'entre otros', r'entre otras',
 r'cabe (?:destacar|mencionar|se[ñn]alar|resaltar|recordar|agregar|subrayar|notar|precisar|apuntar)',
 r'conviene (?:se[ñn]alar|resaltar|destacar|mencionar|precisar)', r'lo anterior (?:permite|muestra|evidencia|refleja|implica)',
 r'es (?:necesario|preciso|importante|fundamental) (?:se[ñn]alar|mencionar|destacar|resaltar|considerar|comprender|entender)',
 r'resulta (?:evidente|claro|importante|fundamental|necesario|pertinente|conveniente)',
 r'm[uú]ltiples', r'diversos autores', r'amplia gama', r'gran potencial', r'gran herramienta',
 r'(?:juega|desempe[ñn]a) un papel', r'herramienta (?:poderosa|fundamental|clave)', r'de gran (?:importancia|relevancia|utilidad)',
 r'entre otros aspectos', r'un sinn[uú]mero de', r'se pueden mencionar', r'como ya se (?:mencion|se[ñn]al)',
 r'tal y como se (?:mencion|se[ñn]al)', r'\betc\.']
BLACKLIST = BLACKLIST + STRICT
OPENERS = [r'sin embargo', r'no obstante', r'asimismo', r'adem[aá]s', r'por lo tanto', r'por ende', r'por consiguiente',
 r'de igual (?:modo|manera|forma)', r'de esta (?:manera|forma)', r'de tal (?:manera|forma)', r'en este sentido',
 r'en ese sentido', r'por su parte', r'por otro lado', r'por otra parte', r'en suma', r'en s[ií]ntesis', r'en resumen',
 r'finalmente', r'por [uú]ltimo', r'cabe \w+', r'con base en lo anterior', r'de acuerdo con lo anterior', r'a su vez']

def sentences(t): return [p for p in re.split(r'(?<=[\.\!\?])\s+', t.strip()) if p.strip()]

def analyze(text):
    low = text.lower()
    bl = sum(len(re.findall(p, low)) for p in BLACKLIST); sym = sum(len(re.findall(p, low)) for p in SYMM)
    en = sum(len(re.findall(p, low)) for p in ENUM); dash = text.count('—')
    lens = [len(s.split()) for s in sentences(text) if s.split()]; words = sum(lens)
    if len(lens) >= 2:
        mean = sum(lens) / len(lens); var = sum((x - mean) ** 2 for x in lens) / len(lens); cv = (var ** 0.5) / mean if mean else 0
    else:
        cv = 0.5
    ss = sentences(text)
    op = sum(1 for s_ in ss if any(re.match(r'\W*' + o + r'\b', s_.lower()) for o in OPENERS))
    patterns = bl + sym + en + dash + op; dens = (patterns / words * 100) if words else 0
    risk = round(min(100.0, min(60.0, dens * 30.0) + (max(0.0, (0.35 - cv) / 0.35) * 40.0 if cv < 0.35 else 0.0)))
    return dict(words=words, bl=bl, sym=sym, en=en, dash=dash, op=op, cv=round(cv, 2), risk=risk, patterns=patterns)

def level(r):
    if r < 20: return 'BAJO'
    if r < 45: return 'MEDIO'
    return 'ALTO'

def markers(body):
    low = body.lower(); found = {}
    for pat in BLACKLIST + SYMM + ENUM:
        for m in re.finditer(pat, low):
            k = m.group(0).strip(); found[k] = found.get(k, 0) + 1
    return [{"t": k, "n": v} for k, v in sorted(found.items(), key=lambda x: -x[1])[:12]]

def main():
    if len(sys.argv) < 2:
        print(json.dumps({"error": "falta el archivo"})); return
    try:
        d = docx.Document(sys.argv[1])
    except Exception as e:
        print(json.dumps({"error": "no se pudo leer el .docx: %s" % e})); return

    sections = []; cur = {'title': 'Preliminares', 'text': []}
    for p in d.paragraphs:
        t = p.text.strip(); st = (p.style.name if p.style is not None else '')
        if st.startswith('Heading') or st.startswith('Título'):
            if cur['text']: sections.append(cur)
            cur = {'title': t or st, 'text': []}
        elif t:
            cur['text'].append(t)
    if cur['text']: sections.append(cur)

    secs = []
    for s in sections:
        if s['title'].strip().lower() == 'preliminares': continue
        body = '\n'.join(s['text'])
        if len(body.split()) < 25: continue
        a = analyze(body); a['title'] = s['title']; a['mk'] = markers(body); secs.append(a)

    # fallback: si el docx no trae Headings, analiza todo el texto como un bloque
    if not secs:
        body = '\n'.join(p.text for p in d.paragraphs if p.text.strip())
        if len(body.split()) >= 25:
            a = analyze(body); a['title'] = 'Documento'; a['mk'] = markers(body); secs.append(a)

    tw = sum(s['words'] for s in secs)
    gr = round(sum(s['risk'] * s['words'] for s in secs) / tw) if tw else 0
    out = {
        "score": gr,
        "level": level(gr),
        "stats": {
            "palabras": tw,
            "secciones": len(secs),
            "muletillas": sum(s['bl'] for s in secs),
            "simetricas": sum(s['sym'] for s in secs),
            "enumeraciones": sum(s['en'] for s in secs),
            "guiones": sum(s['dash'] for s in secs),
        },
        "sections": [
            {"title": s['title'], "risk": s['risk'], "level": level(s['risk']), "words": s['words'],
             "bl": s['bl'], "sym": s['sym'], "en": s['en'], "dash": s['dash'], "cv": s['cv'], "markers": s['mk']}
            for s in secs
        ],
    }
    print(json.dumps(out, ensure_ascii=False))

if __name__ == '__main__':
    main()
